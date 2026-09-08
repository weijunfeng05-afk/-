import asyncio
from functools import lru_cache
from io import BytesIO
import pytest
from cryptography.fernet import Fernet
from docx import Document
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.domain import Requirements, ParsedResume, Evaluation, AppError

REVIEW = {'claimed_ai_depth': '简历自述：未提供 AI 实践深度', 'verification_confidence': '中',
          'ownership_level': '未知', 'strengths': ['有项目经验'], 'risks': [],
          'unknowns': ['个人贡献待核验'], 'must_verify': ['核验个人负责的项目环节'],
          'final_recommendation': '建议 HR 在面试中人工复核具体项目成果。'}


@lru_cache(maxsize=100)
def docx_bytes(text='Alice Chen\nPython backend engineer. Built API services for 5 years.', table=False):
    d = Document()
    for line in text.split('\n'):
        d.add_paragraph(line)
    if table:
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = 'Project experience'
        t.cell(0, 1).text = 'Built a Python API'
    out = BytesIO()
    d.save(out)
    return out.getvalue()


class FixedAdapter:
    calls = []
    on_evaluate = None
    fail = False

    def __init__(self, config):
        self.config = config

    async def test_connection(self):
        return {'ok': True, 'message': '连接成功'}

    async def parse_jd(self, text):
        return Requirements(skills=['Python'], experience=['Backend services'])

    async def parse_resume(self, blocks):
        FixedAdapter.calls.append(('parse', blocks))
        if FixedAdapter.fail:
            raise AppError('model_timeout', '模型超时', retryable=True)
        fact = next((b for b in blocks if 'Python' in b['text']), blocks[-1])
        return ParsedResume.model_validate({'name': 'Alice Chen' if any('Alice Chen' in b['text'] for b in blocks) else '',
            'skills': [{'text': 'Python', 'evidence': [{'text_block_id': fact['id'], 'quote': fact['text']}]}], 'experience': [], 'education': [], 'projects': []})

    async def evaluate_match(self, requirements, parsed, blocks):
        FixedAdapter.calls.append(('evaluate', blocks, parsed))
        if FixedAdapter.on_evaluate:
            FixedAdapter.on_evaluate()
        text = ' '.join(b['text'] for b in blocks)
        fact = next((b for b in blocks if 'Python' in b['text']), blocks[-1])
        evidence = [{'text_block_id': fact['id'], 'quote': fact['text']}]
        unknown = 'UNKNOWN' in text
        return Evaluation.model_validate({'dimensions': [{'dimension': k, 'score': None if unknown else 85, 'reason': '相关原文证据', 'evidence': [] if unknown else evidence} for k in requirements.active()],
            'conditions': [{'condition_id': c.id, 'status': 'unknown' if unknown else 'not_met' if 'NOT_MET' in text else 'met', 'reason': '根据简历判断', 'evidence': [] if unknown else evidence} for c in requirements.conditions],
            'core_information_missing': unknown, 'summary': '依据原文核对岗位要求。',
            'review': REVIEW if requirements.scoring_profile != 'generic' else None})


@pytest.fixture
def app(tmp_path):
    FixedAdapter.calls = []
    FixedAdapter.on_evaluate = None
    FixedAdapter.fail = False
    return create_app(tmp_path, Fernet.generate_key(), worker_enabled=False, model_factory=FixedAdapter)


@pytest.fixture
def client(app):
    with TestClient(app) as c:
        yield c


@pytest.fixture
def ready(client):
    response = client.put('/api/model-config', json={'base_url': 'https://model.example/v1', 'model': 'test-model', 'api_key': 'secret-example-key'})
    assert response.status_code == 200
    job = client.post('/api/jobs', json={'name': 'Python 后端', 'jd_text': 'Python backend role', 'requirements': {'skills': ['Python'], 'conditions': [{'id': 'must-python', 'text': '掌握 Python', 'kind': 'must'}]}, 'confirmed': True}).json()
    return job


def upload(client, job_id, text=None, filename='candidate.docx'):
    return client.post('/api/resumes', data={'job_id': job_id}, files=[('files', (filename, docx_bytes(text) if text else docx_bytes()))]).json()['files'][0]['resume_id']


def drain(app):
    async def run():
        while await app.state.worker.run_once():
            pass
    asyncio.run(run())
