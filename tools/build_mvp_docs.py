"""Create the two MVP documents from the selected System Design template."""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
from lxml import etree as E
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REF = Path('C:/Users/admin/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-system-design/assets/reference.docx')
OUT = ROOT / 'docs'
QA = ROOT / '.qa'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def set_text(p, text):
    rpr = p.find('.//w:rPr', NS)
    rpr = deepcopy(rpr) if rpr is not None else OxmlElement('w:rPr')
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)
    r = OxmlElement('w:r'); r.append(rpr)
    for index,line in enumerate(text.split('\n')):
        if index: r.append(OxmlElement('w:br'))
        t = OxmlElement('w:t'); t.text = line; t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t)
    p.append(r)

class Builder:
    def __init__(self, kind):
        self.kind = kind
        self.src = Document(REF)
        self.root = E.fromstring(ZipFile(REF).read('word/document.xml'))
        self.body = self.root.find('w:body', NS)
        self.sect = deepcopy(self.body.find('w:sectPr', NS))
        children = list(self.body)
        # Keep the source cover, including metadata tables, until its first heading.
        stop = next(i for i,x in enumerate(children) if ''.join(x.itertext()).find('1.  Abstract') >= 0)
        for x in children[stop:]: self.body.remove(x)
        ps = self.body.findall('w:p', NS)
        set_text(ps[8], '简历筛选助手')
        set_text(ps[9], 'MVP ' + kind + ' V1.0')
        # CJK line boxes are taller; remove four blank spacer paragraphs so
        # the complete metadata table remains on the source cover page.
        for p in ps[12:16]: self.body.remove(p)
        tables = self.body.findall('w:tbl', NS)
        self.fill_existing(tables[0], [['版本\nV1.0', '', '使用对象\n单人本机使用', '', '更新日期\n2026年9月7日']])
        self.fill_existing(tables[1], [
            ['文档类型', '第一版 MVP ' + kind],
            ['建设目标', '跑通岗位 JD 与批量简历的匹配和解释流程'],
            ['核心流程', '确认岗位要求 → 上传简历 → 评分分类 → 查看依据'],
            ['适用范围', '单岗位批量匹配，支持 PDF 和 DOCX，接入可配置模型 API'],
        ])
        self.pages = 1

    def fill_existing(self, tbl, rows):
        for tr, vals in zip(tbl.findall('w:tr', NS), rows):
            for tc, text in zip(tr.findall('w:tc', NS), vals):
                ps = tc.findall('w:p', NS)
                # Retain the two-paragraph metadata component where present.
                pieces = text.split('\n')
                for i,p in enumerate(ps): set_text(p, pieces[i] if i < len(pieces) else '')

    def p(self, text, source=22):
        p = deepcopy(self.src.paragraphs[source]._p)
        set_text(p, text)
        ppr = p.find(qn('w:pPr'))
        if ppr is not None:
            if '\n' in text:
                jc=ppr.find(qn('w:jc'))
                if jc is not None: jc.set(qn('w:val'),'left')
            for tag in ('w:numPr','w:pageBreakBefore'):
                x = ppr.find(qn(tag))
                if x is not None: ppr.remove(x)
        self.body.append(p)
        return p

    def page(self, title):
        p = self.p(title, 21)
        pr = p.find(qn('w:pPr'))
        if pr is None: pr = OxmlElement('w:pPr'); p.insert(0,pr)
        br = OxmlElement('w:pageBreakBefore'); pr.append(br)
        self.pages += 1

    def h(self, text): self.p(text, 35)

    def table(self, headers, rows, widths=None):
        source = {2:2,3:5,4:4}[len(headers)]
        tbl = deepcopy(self.src.tables[source]._tbl)
        trs = tbl.findall('w:tr', NS)
        for tr in trs: tbl.remove(tr)
        for i, vals in enumerate([headers] + rows):
            tr = deepcopy(trs[0] if i==0 else trs[1+(i-1)%min(2,len(trs)-1)])
            for tc,text in zip(tr.findall('w:tc',NS), vals):
                ps = tc.findall('w:p',NS)
                set_text(ps[0], text)
                if i == 0:
                    rpr=ps[0].find('w:r/w:rPr',NS)
                    color=rpr.find(qn('w:color'))
                    if color is None: color=OxmlElement('w:color');rpr.append(color)
                    color.attrib.clear();color.set(qn('w:val'),'FFFFFF')
                    bold=OxmlElement('w:b');rpr.append(bold)
                for p in ps[1:]: tc.remove(p)
            tbl.append(tr)
        if widths:
            grid = tbl.find('w:tblGrid',NS)
            for col,w in zip(grid, widths): col.set(qn('w:w'),str(w))
            for tr in tbl.findall('w:tr',NS):
                for tc,w in zip(tr.findall('w:tc',NS),widths):
                    pr=tc.find('w:tcPr',NS)
                    tw=pr.find('w:tcW',NS)
                    if tw is None: tw=OxmlElement('w:tcW');pr.append(tw)
                    tw.set(qn('w:w'),str(w));tw.set(qn('w:type'),'dxa')
        self.body.append(tbl)
        self.p('')

    def save(self):
        self.body.append(self.sect)
        name = '简历筛选助手_MVP_' + self.kind + '_V1.0.docx'
        path = OUT / name
        with ZipFile(REF) as src, ZipFile(path,'w',ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == 'word/document.xml': data=E.tostring(self.root,xml_declaration=True,encoding='UTF-8',standalone=True)
                elif item.filename.startswith('word/footer') and item.filename.endswith('.xml'):
                    data=data.replace(b'[Organization Name] | System Design RFC', ('简历筛选助手 | MVP '+self.kind+' V1.0').encode())
                elif item.filename=='docProps/core.xml':
                    core=E.fromstring(data)
                    for el in core:
                        if E.QName(el).localname in ('creator','lastModifiedBy'): el.text=''
                        if E.QName(el).localname=='title': el.text='简历筛选助手 MVP '+self.kind+' V1.0'
                    data=E.tostring(core,xml_declaration=True,encoding='UTF-8',standalone=True)
                dst.writestr(item,data)
        return {'path':str(path),'expected_pages':self.pages}

def product():
    d=Builder('产品方案')
    d.page('1 产品目标与范围')
    d.p('第一版定位为轻量的 AI 简历筛选助手，面向自行筛选候选人的招聘方。用户上传岗位 JD 与多份简历，补充筛选条件，系统生成岗位匹配评分、分类与原文依据，帮助用户决定优先查看和沟通的候选人。')
    d.p('本版目标是完成可实际使用的筛选流程，并通过人工核对样例判断结果是否有用。评分代表岗位匹配程度，不代表录用概率；系统不自动淘汰或联系候选人。')
    d.table(['纳入第一版','暂不纳入'],[
        ['创建岗位，粘贴或上传 JD，确认提取要求','账号权限、多租户、招聘流程管理'],
        ['一次选择一个岗位，批量上传并匹配简历','跨岗位自动推荐、招聘网站对接'],
        ['硬性条件、加分条件、评分分类与证据','复杂权重编辑、统计报表和消息通知'],
        ['PDF 和 DOCX 文本解析，可配置模型接口','扫描件 OCR、自动面试或候选人联系'],
    ])
    d.h('用户操作流程')
    d.p('创建岗位 → 确认 AI 提取的岗位要求 → 添加额外条件 → 上传简历 → 开始匹配 → 查看排序与详情。')
    d.p('JD 原文和提取后的要求同时保存。用户可以纠正误提取内容；只有确认后的要求才作为匹配依据。同一份简历可分别参与不同岗位匹配。')
    d.h('额外条件')
    d.p('每条额外条件填写一段文字，并选择“必须满足”或“加分项”。例如必须掌握 Python，或有电商项目经验加分。条件针对岗位需要填写，不根据简历中未说明的内容推断。')

    d.page('2 页面与交互')
    d.h('岗位与筛选条件')
    d.p('岗位列表显示岗位名称、简历数量、分析概况和创建时间。新建或编辑岗位时，输入名称与 JD，可上传 PDF、DOCX；点击解析后展示可编辑的技能、经验和其他要求，并添加硬性条件及加分条件。')
    d.h('简历匹配工作台')
    d.p('顶部展示所选岗位、批量上传区域和“开始匹配”。处理中逐份展示等待中、分析中、已完成或失败状态；列表可按分类筛选，并在可排名结果内按总分降序排列。')
    d.table(['列表字段','展示内容'],[
        ['候选人','姓名；未提取到时显示文件名'],
        ['匹配分与分类','总分或暂定分，优先查看、备选、不太匹配、待补充信息'],
        ['硬性条件','满足、不满足或待确认，以及对应数量'],
        ['简要理由与操作','关键匹配点或缺口；查看详情、失败重试、删除'],
    ])
    d.p('候选人详情采用侧边面板，展示基本信息、各维度得分、匹配点、缺口、每条条件判断及简历原文摘录，并提供原文件入口。信息未提供时明确标注，避免使用肯定语气。')
    d.h('大模型配置')
    d.p('配置 API 地址、API Key 和模型名称，支持保存与测试连接。密钥保存后脱敏显示。连接测试同时验证结构化结果是否可用；失败展示可理解的原因，例如鉴权失败、超时或返回格式不符合要求。')
    d.p('岗位要求或额外条件改变后，已有结果标记“需要重新分析”。用户点击重新匹配后生成新结果，避免将旧评分误当作当前结论。')

    d.page('3 评分与分类规则')
    d.table(['维度','权重','判断内容'],[
        ['技能匹配','40%','核心技能及实际使用证据'],
        ['经历相关性','35%','工作和项目经历与岗位的相关程度'],
        ['其他岗位要求','15%','JD 明确要求的学历、资格等'],
        ['额外加分条件','10%','招聘方补充的偏好条件'],
    ],[2000,1100,7124])
    d.p('每个适用维度输出 0 至 100 分、理由和原文依据，由程序按权重计算总分。某维度没有要求时，其权重按比例分配给其他适用维度。要求存在但简历未提供信息时，不能当成“无要求”移除。')
    d.h('分类按顺序判断')
    d.table(['顺序','分类','触发规则'],[
        ['1','不太匹配','至少一条硬性条件有明确证据证明不满足'],
        ['2','待补充信息','没有明确不满足项，但硬性条件或核心信息不足'],
        ['3','优先查看','信息足够且硬性条件满足，总分达到 80 分'],
        ['4','备选','信息足够且硬性条件满足，总分达到 60 分且低于 80 分'],
        ['5','不太匹配','信息足够且总分低于 60 分'],
    ],[750,1950,7524])
    d.p('未设置硬性条件时，该项检查视为通过。信息不足可展示暂定分，并排除在正常排名之外；文件解析失败独立显示“分析失败”，不能作为候选人不匹配的证据。')
    d.h('结果解释原则')
    d.p('每条判断尽量给出可核对的简历摘录。大模型负责理解与判断，程序负责总分和分类。不能用姓名、联系方式等无关信息影响分数，也不能将简历缺失的信息编造为事实。')

    d.page('4 异常处理与验收')
    d.p('无法提取文字的 PDF 提示扫描件暂不支持；损坏文件或模型调用失败显示具体原因，支持单份重试。一份失败不影响其他简历。刷新页面后仍可看到已保存的岗位、进度和结果。')
    d.table(['验收项目','通过条件'],[
        ['岗位与条件','能上传或粘贴 JD，解析、修改、确认要求并保存额外条件'],
        ['简历上传','PDF、DOCX 能解析正文及表格，扫描件和损坏文件有明确提示'],
        ['匹配闭环','批量简历逐份处理，展示进度、分数、分类和原文依据'],
        ['业务规则','覆盖硬性条件不满足、信息缺失、60 分和 80 分边界'],
        ['结果操作','排序、筛选、详情、查看原文件、删除和失败重试可用'],
        ['状态保留','刷新不丢结果；条件修改后旧结果标记需要重新分析'],
        ['模型配置','可保存、脱敏展示、测试连接；密钥不返回完整明文'],
    ])
    d.h('交付顺序')
    d.p('第一步完成模型配置与岗位管理；第二步完成简历上传和解析；第三步完成后台匹配、评分与分类；第四步完成结果详情与完整流程验收。')
    d.p('准备包含明确匹配、明确不满足、信息不足、格式异常的人工核对样例，检查评分依据是否支持结论。用户对结果有争议时应能定位到条件和原文，作为后续调整规则的依据。')
    d.p('每次实现改动补充相关测试，全部测试和验证通过后创建对应 Git commit。第一版运行于本机，后续是否扩展多人使用与部署环境，依据实际使用反馈再决定。')
    return d.save()

def technical():
    d=Builder('技术方案')
    d.page('1 架构与技术选型')
    d.p('采用单机 Web 应用架构：React 前端负责岗位管理、文件上传与结果展示，Python 后端负责文件解析、模型调用和评分，SQLite 保存业务数据与任务状态。本方案是第一版实现设计，尚不代表功能已开发完成。')
    d.table(['模块','技术选择','职责'],[
        ['前端','React + TypeScript + Vite','三个页面、上传、进度轮询、排序筛选与详情'],
        ['后端','Python + FastAPI','业务 API、输入校验、文件服务和模型适配'],
        ['数据与文件','SQLite + 本地目录','保存岗位、简历、任务、结果和原文件'],
        ['任务执行','数据库任务表 + 单执行器','持久化排队，首版顺序执行，失败可重试'],
        ['模型调用','统一适配层','可配置 Base URL、密钥、模型与响应校验'],
    ],[1600,3150,5474])
    d.h('模块交互')
    d.p('浏览器 → 后端 API → 数据库及原文件目录。分析请求先写入任务表并返回任务编号；后台执行器依次完成文本提取、结构化解析、岗位匹配、程序计分，将状态和结果写回数据库。浏览器定时查询任务进度。')
    d.p('模型接口是唯一必需的外部依赖。第一版不引入 Redis、分布式队列、向量数据库或跨岗位检索。原文件和解析记录可复用，匹配结果绑定具体岗位版本。')
    d.h('运行边界')
    d.p('首版单人本机使用，后端仅监听本机回环地址，构建后的前端由后端同源提供。API 服务和一个后台执行器可以作为两个受启动脚本管理的进程；数据库事务保持短小，外部模型请求不占用写事务。')
    d.p('React 官方提供 Vite 搭建方式；FastAPI 支持上传文件；SQLite 适用于单机应用和低写并发使用。版本在开发时选定并通过依赖锁文件固定，避免文档中预先锁定未经验证的版本。')

    d.page('2 文件与模型处理流程')
    d.h('JD 解析')
    d.p('上传或粘贴 JD → 提取文本 → 模型提取岗位要求 → 后端校验结构 → 用户编辑并确认 → 保存岗位要求版本。未确认的解析结果不能直接触发匹配。额外硬性条件和加分条件一并纳入版本。')
    d.h('简历解析')
    d.p('校验文件类型与大小 → 使用随机文件标识保存原文件 → 提取段落与表格文字 → 生成可引用文本块 → 模型抽取技能、工作、项目和教育信息 → 保存结构化内容。文件扩展名与实际可解析内容均需检查。')
    d.p('PDF 保留页码和文本块编号；DOCX 保留段落或表格单元格编号，不承诺固定页码。扫描 PDF、损坏文件、空文本分别报错。超过模型上下文上限时明确失败或分块解析，禁止静默截断。具体文件和文本上限在开发阶段配置并测试。')
    d.h('岗位匹配')
    d.p('加载已确认的岗位版本、额外条件、简历解析版本与评分规则版本 → 调用模型输出维度评分、条件判断及证据 → 后端校验 → 计算总分和分类 → 原子保存结果及完成状态。')
    d.p('同一原文件通过 SHA256 指纹识别，可复用已成功的文本提取和结构化解析。岗位要求变化后仅重做匹配；简历文件、解析规则或模型提取配置改变时使用新解析版本。')
    d.h('模型适配契约')
    d.p('提供 parse_jd、parse_resume、evaluate_match、test_connection 四个内部能力。首版约定一种兼容的聊天补全 HTTP 接口形态；结构化输出能力经连接测试确认后使用，后端始终执行 JSON 和字段校验。')
    d.p('模型输出中每条条件包含 condition_id、status、reason、evidence；status 限定为 met、not_met、unknown。每个引用包含 text_block_id 和原文摘录；维度评分为有限数值 0 至 100 或信息不足时为空。')
    d.p('简历和 JD 作为数据输入，与系统指令分离。其中文字不能改变评分规则、触发工具或要求返回密钥。匹配只发送必要内容，姓名、电话和邮箱不进入评分输入。')

    d.page('3 数据结构与版本')
    d.table(['表','关键字段','约束与用途'],[
        ['jobs','id、name、jd_text、requirements_json、version','保存当前确认要求；修改要求时版本递增'],
        ['resumes','id、file_hash、file_path、text_blocks、parsed_json、parse_version','原文件、引用位置和可复用解析结果'],
        ['match_results','id、job_id、resume_id、input_snapshot、scores、category、evidence','保存岗位快照、总分、理由及各类版本'],
        ['analysis_tasks','id、task_type、input_key、status、attempt、lease_until、error','持久化任务、去重、执行租约及重试记录'],
        ['model_configs','id、base_url、model、encrypted_key、config_version','服务端配置；接口只返回密钥是否设置及掩码'],
    ],[1650,3850,4724])
    d.h('版本与关联')
    d.p('匹配输入快照保存当时的岗位要求与额外条件，结果记录 job_version、parse_version、model、config_version、prompt_version 和 scoring_version。配置版本不包含密钥。即使当前岗位被修改，旧结果仍能说明基于哪些输入得到。')
    d.p('任务 input_key 由任务类型、岗位版本、简历解析版本和规则配置版本组成。同一输入正在等待或执行时返回现有任务；已有成功结果默认复用，用户明确重新分析时创建新的执行记录。')
    d.h('存储与一致性')
    d.p('每份匹配结果和对应任务的完成状态在同一事务提交。模型调用在事务外执行。SQLite 设置合理的繁忙等待；初版仅一个分析执行器，减少并发写冲突。文件写入失败时不生成可分析记录。')
    d.p('删除简历时先停止接收其新任务，再取消等待任务或标记执行任务丢弃结果，删除关联结果和解析记录，最后清理原文件。执行器提交前重新检查记录是否有效，避免删除后重新写回。')

    d.page('4 业务 API 与界面联动')
    d.table(['方法','路径','作用'],[
        ['GET / POST','/api/jobs','查询或创建岗位'],
        ['GET / PATCH','/api/jobs/{id}','查看岗位或修改要求，并更新版本'],
        ['POST','/api/jobs/{id}/parse','创建 JD 解析任务，返回任务编号'],
        ['POST','/api/resumes','批量上传简历，返回逐文件接收结果'],
        ['POST','/api/jobs/{id}/matches','提交 resume_ids，创建或复用匹配任务'],
        ['GET','/api/tasks/{id}','查询状态、处理阶段和脱敏错误'],
        ['POST','/api/tasks/{id}/retry','重试失败或中断任务'],
        ['GET','/api/jobs/{id}/matches','分页查询结果，支持排序与分类筛选'],
        ['GET','/api/matches/{id}','查询维度、条件判断、证据及版本'],
        ['GET / DELETE','/api/resumes/{id}','查看简历记录或删除简历及关联数据'],
        ['GET','/api/resumes/{id}/file','通过服务端标识读取原文件'],
        ['GET / PUT','/api/model-config','读取脱敏配置或保存配置'],
        ['POST','/api/model-config/test','测试鉴权、连接和结构化响应'],
    ],[1800,4200,4224])
    d.p('异步操作返回 202 与 task_id；批量提交返回各简历对应的任务编号。前端建议每 2 秒查询进行中的任务，终态停止轮询。刷新后从岗位和任务接口恢复页面，不依赖浏览器内存保存分析状态。')
    d.p('错误响应使用统一的 code、message、retryable 和 request_id。错误不包含完整密钥、完整简历或服务端堆栈。删除、文件读取和参数校验均通过后端执行，不接受客户端提供的任意磁盘路径。')

    d.page('5 评分与后台任务机制')
    d.h('确定性的总分与分类')
    d.p('模型提供各维度 0 至 100 的得分，程序按技能 40%、经历 35%、其他岗位要求 15%、额外加分 10% 加权。没有要求的维度不适用，剩余权重归一化；全部不适用时拒绝开始匹配并提示完善岗位要求。')
    d.p('总分保存并展示到小数点后一位，使用同一数值比较 60 分和 80 分阈值。要求存在但证据缺失时标记 unknown，不移除维度或当作零分。关键条件或核心信息不足时总分可为空；如果展示暂定分，必须带 provisional 标记且不参与正常排名。')
    d.p('分类顺序固定：明确硬性条件不满足 → 不太匹配；否则核心信息或硬性条件不足 → 待补充信息；否则按总分划分优先查看、备选、不太匹配。未设置硬性条件时，该项检查通过。模型不能直接覆盖最终分类。')
    d.h('输出与证据校验')
    d.p('校验必填字段、枚举、分数范围、重复或遗漏的条件 ID。引用须匹配提取文本中的实际片段；仅做空白规范化，不能用相似文本替代原文。引用不存在时本次结果无效，不能直接显示为已验证结论。校验能确认引用存在，不能证明语义判断一定正确。')
    d.h('任务状态与故障恢复')
    d.p('任务状态：queued → running → succeeded 或 failed；重启后租约已过期的 running 任务标记 interrupted，供用户重试。执行阶段另记为文件提取、结构化解析、岗位匹配和结果校验，便于定位失败。')
    d.p('执行器通过短事务领取任务并设置租约，执行中更新心跳；提交时校验执行令牌，避免过期执行覆盖新结果。外部模型请求无法保证恰好一次，超时后重试可能增加调用费用，但数据库结果可以避免重复提交。')
    d.p('网络超时、限流和服务端临时错误建议最多自动重试 2 次并退避；鉴权错误和不支持的文件不自动重试。结构化输出错误可修正重试 1 次，仍不合格则失败。具体超时与限额作为服务端配置，不把任务永久留在分析中。')

    d.page('6 部署与验证计划')
    d.h('配置与运行')
    d.p('构建 React 静态页面，由 FastAPI 提供同源页面和 /api 接口；本机启动脚本同时管理 API 与后台执行器。关闭应用时停止领取新任务，未完成任务依靠租约恢复。初版只服务本机；开放外部访问前另行设计鉴权和部署。')
    d.p('模型密钥由服务端加密保存，加密主密钥由本机环境配置提供，与数据库分开且不提交 Git。日志仅记录任务编号、阶段、耗时和错误类型，不记录完整简历、密钥或原始模型请求。备份应同时覆盖数据库和原文件，并单独管理密钥恢复方式。')
    d.table(['验证层','覆盖内容','通过条件'],[
        ['规则单元测试','权重归一化、空要求、缺失信息、硬性条件、60/80 分边界','计算和分类符合产品规则，结果不自相矛盾'],
        ['解析与契约测试','PDF 正文、DOCX 表格、扫描件、损坏文件、非法 JSON、伪造引用','有效内容无关键遗漏，异常明确失败'],
        ['任务与接口集成','重复提交、部分失败、超时、重启、删除期间完成任务','状态可恢复，无重复或过期结果写回'],
        ['完整流程验收','创建岗位、上传多份简历、分析、筛选、详情、配置测试','真实模型至少跑通一组人工核对样例'],
    ],[1750,4600,3874])
    d.p('日常自动测试使用固定模型响应覆盖可重复规则；真实模型验收核对证据与结论，不要求两次模型输出逐字相同。每次实现改动更新相关测试，通过全部测试和验证后创建对应 Git commit。')
    d.h('实施顺序')
    d.p('M1 模型配置与岗位确认；M2 文件解析与任务持久化；M3 匹配契约、规则计分和结果详情；M4 故障恢复、人工样例核对及本机启动交付。文件大小、模型上下文上限、超时值和依赖版本在实现时配置并验证。')
    d.h('技术参考')
    d.p('React 构建说明 https://react.dev/learn/build-a-react-app-from-scratch\nFastAPI 文件上传 https://fastapi.tiangolo.com/tutorial/request-files/\nSQLite 使用场景 https://www.sqlite.org/whentouse.html')
    return d.save()

if __name__=='__main__':
    OUT.mkdir(exist_ok=True); QA.mkdir(exist_ok=True)
    results=[product(),technical()]
    with ZipFile(REF) as z:
        inventory={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
    report={'reference':str(REF),'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'parts':inventory,'outputs':results}
    (QA/'manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(results,ensure_ascii=False,indent=2))
